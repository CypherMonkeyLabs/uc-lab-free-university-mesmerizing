#!/usr/bin/env python3
"""Write UC university HTML, SVG, sheets, and office exports."""
from __future__ import annotations

import html
import json
import re
import zipfile
from pathlib import Path

from build_uc_round2 import (
    hero_stack_svg,
    CONTACT,
    ORG,
    SECTIONS,
    TRACKS,
    WEB,
    arch_svg,
    build_diagrams,
    dl_buttons,
    esc,
    flow_svg,
    footer,
    rooms,
    write_page,
)
from go_deeper import (
    CATALOG,
    CHEAT_KEYS,
    CORE,
    ROOM_KEYS,
    SHEET_EXTRA,
    TEMPLATE_SHEET,
    union,
    add_docx_sources,
    add_pptx_sources,
    append_pdf_sources,
    entries,
    html_block,
    keys_for_diagram,
    keys_for_sheet,
    keys_for_track,
    markdown_block,
    mirror_sources,
    patch_flagship_cheats,
    plain_lines,
    room_markdown,
    sheet_markdown,
    write_campus_bundle,
    write_json,
    write_text,
)

ROOT = Path(__file__).resolve().parent.parent


def diagram_svg(item: dict) -> str:
    if item["kind"] == "flow":
        return flow_svg(item)
    return arch_svg(item)


def room_body(room: dict, pre: str) -> str:
    bullets = "".join(f"<li>{esc(item)}</li>" for item in room["five"])
    questions = []
    for index, (prompt, why, choices, answer) in enumerate(room["quiz"], start=1):
        buttons = "".join(
            f'<button type="button" class="choice ghost" data-choice="{esc(choice)}">{esc(choice)}</button>'
            for choice in choices
        )
        questions.append(
            f'<div class="quiz" data-quiz data-answer="{esc(answer)}"><p><strong>{index}.</strong> {esc(prompt)}</p>'
            f"{buttons}<p class=\"quiz-note\" hidden>{esc(why)}</p></div>"
        )
    sources = "".join(f'<li><a href="{href}">{esc(title)}</a></li>' for title, href in room["sources"])
    source_items = entries(ROOM_KEYS[room["slug"]])
    mirror_sources(ROOT / "rooms" / "sources" / room["slug"] / "sources.json", room["slug"], source_items)
    deeper = html_block(
        source_items,
        f"sources/{room['slug']}/sources.json",
        md_href=f"{pre}downloads/rooms/{room['slug']}.md",
    )
    glossary = "".join(f"<li><strong>{esc(term)}</strong> — {esc(meaning)}</li>" for term, meaning in room["glossary"])
    confirm = f'<p class="confirm">Confirm before relying. {esc(room["confirm"])}</p>' if room.get("confirm") else ""
    sheet = room["sheet"]
    exports = dl_buttons(pre, room["diagram"], sheet).replace(
        'class="room-downloads"', 'class="room-downloads section-dl"', 1
    )
    return f"""
<div class="wrap">
<p class="kicker">{esc(room["track"])} · {esc(room["id"])}</p>
<h1>{esc(room["title"])}</h1>
{dl_buttons(pre, room["diagram"], sheet)}
<section id="five" class="five"><h2>Learn it in 5 minutes</h2><ul>{bullets}</ul>{exports}</section>
<section id="visual"><h2>Visual</h2><div class="diagram-frame"><img src="{pre}assets/diagrams/university/{room["diagram"]}.svg" alt="{esc(room["title"])} diagram"></div>{exports}</section>
<section id="core"><h2>What is going on</h2><p>{esc(room["core"])}</p>{confirm}{exports}</section>
<section id="example" class="worked"><h2>Worked example</h2><p>{esc(room["example"])}</p>{exports}</section>
<section id="quiz"><h2>Quiz</h2>{''.join(questions)}{exports}</section>
<section id="eco" class="eco-panel"><h2>Ecosystem</h2>
<p>{esc(room["eco"])}</p>
<ul>{glossary}</ul>
<p><strong>Config.</strong> {esc(room["config"])}</p>
<p><strong>Troubleshooting.</strong> {esc(room["trouble"])}</p>
<p><strong>Bridge.</strong> {esc(room["bridge"])}</p>
<p>Mind map, sister technologies, and future hooks live on the <a href="{pre}ecosystem.html">ecosystem page</a>.</p>
{exports}
</section>
<section id="sources" class="sources"><h2>Sources</h2><ul>{sources}</ul>
<p>Product screens move. Confirm the guide for your release before you change a live system.</p>{exports}{deeper}</section>
</div>
"""


def page_wrap(inner: str) -> str:
    return f'<div class="wrap">{inner}</div>'


def emit_diagrams(items: list[dict]) -> None:
    svg_dir = ROOT / "assets" / "diagrams" / "university"
    html_dir = ROOT / "downloads" / "diagrams"
    svg_dir.mkdir(parents=True, exist_ok=True)
    html_dir.mkdir(parents=True, exist_ok=True)
    room_rows = rooms()
    for item in items:
        svg = diagram_svg(item)
        (svg_dir / f"{item['slug']}.svg").write_text(svg, encoding="utf-8")
        source_items = entries(keys_for_diagram(room_rows, item["slug"]))
        write_json(f"diagrams/{item['slug']}.json", item["slug"], source_items)
        mirror_sources(html_dir / "sources" / item["slug"] / "sources.json", item["slug"], source_items)
        deeper = html_block(
            source_items,
            f"sources/{item['slug']}/sources.json",
            dark=True,
            md_href=f"{item['slug']}.md",
        )
        standalone = (
            "<!DOCTYPE html><html lang=\"en\"><head><meta charset=\"utf-8\">"
            f"<title>{esc(item['title'])}</title>"
            "<link rel=\"stylesheet\" href=\"../../assets/css/footer.css\">"
            "<style>body{margin:0;background:#0b1020;color:#e0e7ff;font-family:system-ui,sans-serif}svg{width:100%;height:auto}.go-deeper{margin:16px}</style>"
            f"</head><body><h1 style=\"padding:12px 16px\">{esc(item['title'])}</h1>{svg}{deeper}{footer()}</body></html>"
        )
        (html_dir / f"{item['slug']}.html").write_text(standalone, encoding="utf-8")
        json_href = f"sources/{item['slug']}/sources.json"
        write_text(
            html_dir / f"{item['slug']}.md",
            f"# {item['title']}\n\nSection export for this drawing.\n\n{markdown_block(source_items, json_href)}",
        )


def emit_rooms(room_rows: list[dict]) -> None:
    for room in room_rows:
        source_items = entries(ROOM_KEYS[room["slug"]])
        write_json(f"{room['slug']}.json", room["slug"], source_items)
        mirror_sources(ROOT / "rooms" / "sources" / room["slug"] / "sources.json", room["slug"], source_items)
        rel = f"rooms/{room['slug']}.html"
        write_page(rel, room["title"], room["five"][0], "learn", room_body(room, "../"), room["track"])
        write_text(
            ROOT / "downloads" / "rooms" / f"{room['slug']}.md",
            room_markdown(room, source_items, f"../../rooms/sources/{room['slug']}/sources.json"),
        )


def emit_tracks(room_rows: list[dict]) -> None:
    for tid, name, blurb, color in TRACKS:
        cards = []
        for room in room_rows:
            if room["track"] != tid:
                continue
            cards.append(
                f'<a class="room-card" href="../rooms/{room["slug"]}.html"><h2>{esc(room["title"])}</h2><p>{esc(room["five"][0])}</p></a>'
            )
        source_items = entries(keys_for_track(room_rows, tid))
        write_json(f"tracks/{tid}.json", tid, source_items)
        mirror_sources(ROOT / "tracks" / "sources" / tid / "sources.json", tid, source_items)
        deeper = html_block(
            source_items,
            f"sources/{tid}/sources.json",
            md_href=f"../downloads/tracks/{tid}.md",
        )
        body = page_wrap(
            f'<p class="kicker">{esc(tid)}</p><h1>{esc(name)}</h1><p class="lede">{esc(blurb)}</p>'
            f'<div class="room-downloads"><a class="dl-btn" href="../downloads/sheets/{sheet_for(tid)}.html">Cheat sheet</a>'
            f'<a class="dl-btn" href="../downloads/tracks/{tid}.md">Markdown</a></div>'
            f'<div class="card-grid">{"".join(cards)}</div>{deeper}'
        )
        write_page(f"tracks/{tid}.html", name, blurb, "learn", body, tid)
        room_lines = "\n".join(f"- {room['title']}" for room in room_rows if room["track"] == tid)
        write_text(
            ROOT / "downloads" / "tracks" / f"{tid}.md",
            f"# {name}\n\n{blurb}\n\n{room_lines}\n\n{markdown_block(source_items, f'../../tracks/sources/{tid}/sources.json')}",
        )


def sheet_for(tid: str) -> str:
    return {
        "T1": "sip",
        "T2": "dial-plan",
        "T3": "cube",
        "T4": "expressway",
        "T5": "qos",
        "T6": "ucce",
        "T7": "migration",
        "T8": "trouble",
        "T9": "e911",
        "T10": "recording",
        "T11": "bridges",
    }[tid]


SHEETS = {
    "sip": ("SIP in one page", ["A transaction is not a dialog.", "180 is alerting. 200 accepts. ACK finishes a 2xx.", "REGISTER stores a Contact.", "Early offer puts SDP in the INVITE.", "PRACK answers a reliable 183.", "RTP may skip the proxy.", "Source: RFC 3261, RFC 3264, RFC 3262."]),
    "dial-plan": ("CUCM dial plan", ["A CSS is an ordered list of partitions.", "First match in that order wins.", "Route pattern, route list, route group, gateway.", "Called and calling numbers are different fields.", "A trunk in service can still have an empty CSS.", "Source: Cisco Unified Communications Manager documentation."]),
    "cube": ("CUBE, dial peers, SRST", ["CUBE is a B2BUA with two dialogs.", "Confirm inbound match order for your IOS release.", "Longest match, then preference, on outbound peers.", "SRST is a subset dial plan, not a clone of CUCM.", "A local gateway toward Webex Calling is a different peer.", "Source: Cisco CUBE and SRST documentation. Confirm call caps."]),
    "expressway": ("Expressway and MRA", ["C inside, E at the edge, traversal zone between them.", "MRA still registers to CUCM.", "Do not copy a port list from memory.", "Signaling up does not prove media.", "Source: Cisco Expressway deployment guides."]),
    "qos": ("Voice QoS", ["EF is DSCP 46.", "Cisco signaling designs often use CS3. RFC 4594 points signaling at CS5.", "G.711 at 20 ms is about 80 kbit/s at layer 3.", "LLQ schedules bandwidth. It does not create it.", "CUCM locations do not mark packets.", "G.114 is the usual 150 ms planning citation. Confirm the current text."]),
    "ucce": ("UCCE bridge", ["Ingress, script, queue, and agent are different stages.", "A label must be a number CUCM can route.", "Queue audio is not an agent ringing.", "Webex Contact Center is a different product.", "Source: Cisco UCCE design guides."]),
    "migration": ("Webex Calling migration", ["Calling and Dedicated Instance are different offers.", "Dual-run a number you can return.", "Save the old route before you edit it.", "SRST does not automatically cover Webex apps.", "Confirm survivability features on your subscription."]),
    "trouble": ("Troubleshooting habit", ["One hypothesis. One change.", "Read one SIP leg at a time.", "A 200 OK does not prove two-way audio.", "Compare SDP addresses with the packets you see.", "Revert if the symptom does not move."]),
    "e911": ("Emergency calling note", ["Study notes, not legal advice and not an emergency plan.", "The FCC MLTS page describes direct 911 dialing, notification, and dispatchable location for covered US systems.", "Cisco Emergency Responder maps a switch port to an ERL and an ELIN. Confirm Release 15 before you build.", "A PSAP callback is a separate call toward that ELIN.", "Rules and carriers outside the pages this lab opened stay on a confirm row."]),
    "recording": ("Call recording note", ["Phone-based recording forks from the phone. Network-based recording can fork from the phone or the gateway.", "SIPREC is a separate recording session. RFC 7866 is the protocol. RFC 7865 is the metadata.", "CUBE's IOS XE 17.6 guide calls CUBE the session recording client when SIPREC is configured.", "A quality score is not the audio file.", "Consent and retention were not opened. Confirm the rule where you operate."]),
    "bridges": ("Open telephony bridges", ["A SIP trunk joins two dial plans. It does not merge them.", "SIPconnect is the SIP Forum recommendation for provider peering.", "Asterisk and FreeSWITCH are separate PBXs with their own dialplans.", "WebRTC is a browser API plus the protocol suite in RFC 8825. SIP over WebSocket is RFC 7118.", "A CPaaS vendor matrix was not opened. Confirm a vendor on that vendor's page."]),
    "bandwidth": ("Bandwidth sketch", ["G.711 20 ms payload is 160 bytes.", "IP, UDP, and RTP add 40 bytes.", "Layer 3 is about 80 kbit/s.", "An 18-byte Ethernet header makes about 87.2 kbit/s.", "Size the priority queue from admitted calls.", "Formula: (payload + 40 + L2) x 8 x (1000 / ptime) / 1000."]),
    "ports-note": ("Ports, without a fake matrix", ["SIP is often 5060 and 5061. That is not a firewall design.", "Media ranges differ for phones, CUBE, and Expressway.", "Open the table in the guide for your version.", "Write source, destination, and direction.", "Confirm before relying."]),
    "evidence": ("Evidence pack", ["Time, calling number, called number.", "Hypothesis in one sentence.", "The ladder or the trace.", "The single change.", "The result, including a revert."]),
}


def emit_sheets() -> None:
    folder = ROOT / "downloads" / "sheets"
    folder.mkdir(parents=True, exist_ok=True)
    room_rows = rooms()
    for slug, (title, lines) in SHEETS.items():
        source_items = entries(keys_for_sheet(room_rows, slug))
        write_json(f"sheets/{slug}.json", slug, source_items)
        mirror_sources(folder / "sources" / slug / "sources.json", slug, source_items)
        deeper = html_block(source_items, f"sources/{slug}/sources.json", md_href=f"{slug}.md")
        items = "".join(f"<li>{esc(line)}</li>" for line in lines)
        doc = (
            "<!DOCTYPE html><html lang=\"en\"><head><meta charset=\"utf-8\">"
            f"<title>{esc(title)}</title><link rel=\"stylesheet\" href=\"../../assets/css/footer.css\">"
            "<style>body{font-family:system-ui,sans-serif;margin:24px;background:#f8fafc;color:#0f172a}ul{font-size:1.15rem;line-height:1.5}</style>"
            f"</head><body><h1>{esc(title)}</h1><ul>{items}</ul>{deeper}{footer()}</body></html>"
        )
        (folder / f"{slug}.html").write_text(doc, encoding="utf-8")
        write_text(folder / f"{slug}.md", sheet_markdown(title, lines, source_items, f"sources/{slug}/sources.json"))


def study_primer() -> str:
    """Cited learner text for the front door. Counts stay in the prose, not in a widget."""
    return """
<h2>Read this before the first room</h2>
<p>A phone call in this lab is two jobs that share a clock. Signaling asks for the session and names the parties. Media is the voice, carried as RTP. <a href="https://www.rfc-editor.org/rfc/rfc3261">RFC 3261</a> defines the Session Initiation Protocol that carries the request. <a href="https://www.rfc-editor.org/rfc/rfc3550">RFC 3550</a> defines RTP. If you mix those jobs, you will chase a dial plan while the packets never arrive, or you will chase a codec while the request never left the phone.</p>
<p>Start with one leg. An INVITE asks for a session. A 100 Trying means a hop received it. A 180 Ringing means alerting. A 200 OK means the far end accepted. The ACK completes a successful INVITE. A transaction is that request together with its responses. A dialog is the longer relationship, named by a Call-ID and the tags on From and To. A 100 has no To tag, so it is not yet a dialog. BYE ends the dialog you already have. CUCM and CUBE are often back-to-back user agents, so one human call can be two Call-IDs. Count the dialogs before you decide the call failed.</p>
<p>SDP is not a second protocol you can skip. <a href="https://www.rfc-editor.org/rfc/rfc4566">RFC 4566</a> describes the session description. <a href="https://www.rfc-editor.org/rfc/rfc3264">RFC 3264</a> is the offer and the answer. Early offer puts the offer in the INVITE. Delayed offer puts it later, often in the 200, with the answer in the ACK. The answer may drop codecs. It does not invent a codec the offer never listed. Read the m= lines and the direction attributes. sendonly against sendonly is a quiet call with perfect signaling.</p>
<p>REGISTER is not a call. It publishes a Contact for an address-of-record, and the binding expires. A 401 is a challenge. The next REGISTER needs a new CSeq. A phone can be registered and still have no media path. That distinction matters again at the edge. Mobile and Remote Access still ends as a registration on the call agent. The Expressway pair is the path. The port list belongs to the guide for the release you run. This page does not invent one.</p>
<p>Early media is the tone or the announcement you hear before answer. A 183 Session Progress can carry SDP. When that response is reliable, <a href="https://www.rfc-editor.org/rfc/rfc3262">RFC 3262</a> expects a PRACK, and PRACK is its own transaction. Hearing ringback does not mean the called party answered.</p>
<p>The dial plan is who may reach which pattern. On CUCM a partition is a bucket of patterns and a calling search space is an ordered list of those buckets. First match in that order wins. A route pattern matches digits, a route list orders route groups, and a route group holds trunks or gateways. Called, calling, and redirecting numbers are different fields, so a mask can repair caller identity or break it. Menu names move between releases. When a Cisco guide was not opened for this lab, the sources block says confirm and shows no address. Use the guide for the version you run before you change a live system.</p>
<p>CUBE is a border element that terminates one SIP dialog and starts another. Inbound match and outbound match are different questions. SRST is the smaller dial plan a phone uses when it cannot reach the call agent. It is not a copy of the campus dial plan. A local gateway toward Webex Calling is another peer. The local-gateway articles on Webex Help are the ones this lab opened, including <a href="https://help.webex.com/en-us/article/t9xctu/Get-started-with-Local-Gateway">Get started with Local Gateway</a>.</p>
<p>Voice quality is a budget, not a slogan. <a href="https://www.rfc-editor.org/rfc/rfc3246">RFC 3246</a> defines expedited forwarding, the behavior people mark as DSCP 46 for voice RTP. <a href="https://www.rfc-editor.org/rfc/rfc4594">RFC 4594</a> is the service-class map, and it does not always match a vendor signaling mark. <a href="https://www.itu.int/rec/T-REC-G.114">ITU-T G.114</a> is the usual citation for one-way delay planning, often remembered as 150 milliseconds. <a href="https://www.itu.int/rec/T-REC-G.711">ITU-T G.711</a> is the pulse-code codec behind the 64 kbit/s payload. On the wire, a 20 millisecond G.711 packet plus the IP, UDP, and RTP headers is about 80 kbit/s at layer 3, and a little more with an Ethernet header. A priority queue schedules that class. It does not create bandwidth the link does not have. A location on the call agent is a call-admission budget between sites. It does not, by itself, mark the packets.</p>
<p>Contact center is a later stage of the same call. Ingress, the routing script, the queue treatment, and the agent endpoint are different objects. Queue audio is not an agent ringing. Webex Contact Center and Unified Contact Center Enterprise are different products. A label still has to be a number the call agent can route.</p>
<p>A move toward Webex Calling is a number you can give back. Multi-tenant Calling and Dedicated Instance are different offers. <a href="https://help.webex.com/en-us/article/8jda92/Webex-Calling">Get started with Webex Calling</a> is the overview this lab opened on the Webex Help Center. Dual-run one number, save the old route, and rehearse the rollback before you move the next one. Survivability features depend on the subscription. If a survivability article was not opened, the sources row says confirm and carries no address.</p>
<p>Emergency calling is a later requirement on the same phone, and it is not a lab trick. The FCC page this lab opened, <a href="https://www.fcc.gov/mlts-911-requirements">Kari's Law and RAY BAUM'S Act for multi-line telephone systems</a>, describes direct 911 dialing, a notification that must not delay the call, and dispatchable location for covered systems in the United States. <a href="https://www.cisco.com/c/en/us/support/unified-communications/emergency-responder/series.html">Cisco Emergency Responder</a> is the product that maps a discovered switch port to an emergency response location and an ELIN. The Release 15 planning chapter says phones on unknown switches are unlocated. This is study material, not legal advice. A rule this lab did not open stays on a confirm row with no address.</p>
<p>Call recording is a copy of the media, started on purpose. The CUCM 15 feature guide says a phone can fork two streams, or a gateway can, and the built-in bridge has to be on when the phone is the source. <a href="https://www.rfc-editor.org/rfc/rfc7866">RFC 7866</a> is the session recording protocol. <a href="https://www.rfc-editor.org/rfc/rfc7245">RFC 7245</a> is the architecture, and it says some environments record for compliance. It does not name your statute. A quality score is not the file.</p>
<p>An open PBX or a browser is another peer. <a href="https://en.wikipedia.org/wiki/SIP_trunking">SIP trunking</a> joins a provider and a SIP PBX. Asterisk and FreeSWITCH keep their own dialplans. <a href="https://www.rfc-editor.org/rfc/rfc8825">RFC 8825</a> is the WebRTC protocol overview, and <a href="https://www.rfc-editor.org/rfc/rfc7118">RFC 7118</a> carries SIP over a WebSocket. A browser does not become a CUCM phone because both can carry voice.</p>
<p>When something breaks, write one hypothesis and make one change. A 200 OK does not prove two-way audio. Compare the addresses in SDP with the packets you actually see. If the symptom does not move, put the change back. The tracks below are that sequence: SIP, the dial plan, CUBE and SRST, Expressway, voice QoS, the contact-center bridge, a careful migration, the troubleshooting habit, emergency calling, recording, and the open-telephony bridge.</p>
"""


def emit_home(room_rows: list[dict], items: list[dict]) -> None:
    tiles = []
    for tid, name, blurb, color in TRACKS:
        count = sum(1 for room in room_rows if room["track"] == tid)
        tiles.append(
            f'<a class="track-tile" style="--card:{color}" href="tracks/{tid}.html"><h2>{esc(name)}</h2><p>{esc(blurb)}</p><p>{count} rooms</p></a>'
        )
    flows = sum(1 for item in items if item["kind"] == "flow")
    home_items = entries(CORE)
    mirror_sources(ROOT / "sources" / "home" / "sources.json", "home", home_items)
    write_json("home.json", "home", home_items)
    deeper = html_block(home_items, "sources/home/sources.json")
    body = f"""
<section class="wrap hero-band">
<div>
<p class="kicker">CUCM · SIP · gateways · contact center</p>
<h1>Voice, then the dial plan, then the edge.</h1>
<p class="lede">A free study lab for Unified Communications: SIP, CUCM, CUBE, SRST, Expressway, voice QoS, UCCE, a careful move toward Webex Calling, emergency calling, call recording, and bridges to open telephony. Figures that depend on a release are marked confirm before relying. Emergency and recording pages are study notes, not legal advice.</p>
<div class="actions">
<a class="btn" href="start.html">Start here</a>
<a class="ghost" href="diagrams.html">Diagrams</a>
<a class="ghost" href="downloads.html">Downloads</a>
<a class="ghost" href="campus/uc-lab-free-university.html">Deep campus</a>
</div>
</div>
<div class="diagram-frame hero-diagram">{hero_stack_svg()}</div>
</section>
<section class="wrap">
<div class="stat-row" aria-label="Counts from this repository">
<div class="stat"><b data-count="{len(room_rows)}">{len(room_rows)}</b><span>rooms</span></div>
<div class="stat"><b data-count="{len(items)}">{len(items)}</b><span>diagrams</span></div>
<div class="stat"><b data-count="{flows}">{flows}</b><span>flows</span></div>
<div class="stat"><b data-count="{len(TRACKS)}">{len(TRACKS)}</b><span>tracks</span></div>
</div>
<h2>{len(TRACKS)} paths</h2>
<div class="track-map">{''.join(tiles)}</div>
<p class="lede">The older flagship studio remains at <a href="flagship/index.html">flagship/index.html</a>. The 580-section campus remains beside it as a comparison library: Teams Phone, Genesys, Five9, NICE, Amazon Connect, and the other names there are study references, not this course and not endorsements. This front door is the round-2 Cisco UC course.</p>
</section>
<section class="wrap study-primer" id="primer">
{study_primer()}
{deeper}
</section>
"""
    write_page("index.html", "UC Lab Free University", "SIP, CUCM, CUBE, SRST, Expressway, QoS, UCCE, and Webex Calling migration.", "home", body)


def emit_start(room_rows: list[dict]) -> None:
    ten = room_rows[:4]
    links = "".join(f'<li><a href="rooms/{room["slug"]}.html">{esc(room["title"])}</a> — {esc(room["five"][0])}</li>' for room in ten)
    full = "".join(
        f'<li><a href="tracks/{tid}.html">{esc(name)}</a> — {esc(blurb)}</li>' for tid, name, blurb, _color in TRACKS
    )
    start_items = entries(["rfc3261", "rfc3264", "rfc3550", "rfc4566", "rfc3262", "cucm", "wiki-sip"])
    mirror_sources(ROOT / "sources" / "start" / "sources.json", "start", start_items)
    write_json("start.json", "start", start_items)
    deeper = html_block(start_items, "sources/start/sources.json")
    body = page_wrap(
        f"""
<h1>Start here</h1>
<section id="ten"><h2>Finish in 10 minutes</h2>
<p>Read four SIP rooms. You will be able to tell a transaction from a dialog and a 180 from a 200.</p>
<ol>{links}</ol>
<div class="room-downloads"><a class="dl-btn" href="downloads/sheets/sip.html">SIP sheet</a></div>
</section>
<section id="path"><h2>Full course</h2>
<p>Fundamentals, the dial plan, the gateway, the edge, QoS, the contact-center bridge, migration, the troubleshooting habit, then emergency calling, recording, and open-telephony bridges.</p>
<ol>{full}</ol>
</section>
<section><h2>Certifications</h2>
<p>Collaboration exam titles change. The Cisco training catalog returned HTTP 403 to the checker on 2026-10, so this page does not link it. Confirm the current title on the vendor site before you pay. This lab is study material, not an exam voucher.</p>
</section>
{deeper}
"""
    )
    write_page("start.html", "Start here", "Ten minutes or the full UC course.", "learn", body)


def emit_diagram_indexes(items: list[dict]) -> None:
    cards = []
    flows = []
    for item in items:
        cards.append(
            f'<article class="room-card"><h2>{esc(item["title"])}</h2><div class="diagram-frame"><img src="assets/diagrams/university/{item["slug"]}.svg" alt="{esc(item["title"])}"></div>'
            f'<div class="room-downloads"><a class="dl-btn" href="assets/diagrams/university/{item["slug"]}.svg" download>SVG</a>'
            f'<a class="dl-btn" href="downloads/diagrams/{item["slug"]}.png" download>PNG</a>'
            f'<a class="dl-btn" href="downloads/diagrams/{item["slug"]}.html" download>HTML</a></div></article>'
        )
        if item["kind"] == "flow":
            flows.append(cards[-1])
    diagram_sources = entries(CORE)
    diagram_block = html_block(diagram_sources, "sources/diagrams/sources.json")
    flow_block = html_block(diagram_sources, "sources/flows/sources.json")
    mirror_sources(ROOT / "sources" / "diagrams" / "sources.json", "diagrams", diagram_sources)
    mirror_sources(ROOT / "sources" / "flows" / "sources.json", "flows", diagram_sources)
    write_json("diagrams.json", "diagrams", diagram_sources)
    write_json("flows.json", "flows", diagram_sources)
    write_page("diagrams.html", "Diagrams", "UC diagrams as SVG, PNG, and HTML.", "diagrams", page_wrap(f"<h1>Diagrams</h1><p>{len(items)} drawings. Each export has its own source list.</p><div class=\"card-grid\">{''.join(cards)}</div>{diagram_block}"))
    write_page("flows.html", "Flows", "Animated UC call flows.", "flows", page_wrap(f"<h1>Flows</h1><p>{len(flows)} ladders with a moving packet. Reduced motion keeps every step visible in the SVG.</p><div class=\"card-grid\">{''.join(flows)}</div>{flow_block}"))


def emit_charts(room_rows: list[dict], items: list[dict]) -> None:
    bars = []
    for tid, name, _blurb, color in TRACKS:
        count = sum(1 for room in room_rows if room["track"] == tid)
        width = count * 18
        bars.append(
            f'<div data-chart="rooms-{tid}"><span>{esc(name)} · {count}</span><div class="chart-bar" style="width:{width}px;background:{color}"></div></div>'
        )
    flows = sum(1 for item in items if item["kind"] == "flow")
    arches = len(items) - flows
    extra = f"""
<div data-chart="flows"><span>Animated flows · {flows}</span><div class="chart-bar" style="width:{flows * 8}px"></div></div>
<div data-chart="arches"><span>Still diagrams · {arches}</span><div class="chart-bar" style="width:{arches * 8}px"></div></div>
<div data-chart="g711"><span>G.711 layer 3 · 80 kbit/s</span><div class="chart-bar" style="width:160px"></div></div>
<div data-chart="g711-l2"><span>G.711 plus Ethernet · 87.2 kbit/s</span><div class="chart-bar" style="width:174px"></div></div>
"""
    chart_sources = entries(SHEET_EXTRA["bandwidth"])
    chart_block = html_block(chart_sources, "sources/charts/sources.json")
    mirror_sources(ROOT / "sources" / "charts" / "sources.json", "charts", chart_sources)
    write_json("charts.json", "charts", chart_sources)
    body = page_wrap("<h1>Charts</h1><p>Room counts come from this repository. The G.711 bars use the bandwidth sheet sources. No traffic and no revenue.</p>" + "".join(bars) + extra + chart_block)
    write_page("charts.html", "Charts", "Counts from this UC lab.", "diagrams", body)


def emit_play() -> None:
    play_sources = entries(["rfc3261", "rfc3246", "rfc4594", "cube", "cucm"])
    play_block = html_block(play_sources, "sources/play/sources.json")
    mirror_sources(ROOT / "sources" / "play" / "sources.json", "play", play_sources)
    write_json("play.json", "play", play_sources)
    body = page_wrap(
        f"""
<h1>Play</h1>
<p>Three checks. Nothing is stored on a server.</p>
<div class="quiz" data-quiz data-answer="Alerting, not answer"><p>What is 180 Ringing?</p>
<button type="button" class="choice ghost" data-choice="Alerting, not answer">Alerting, not answer</button>
<button type="button" class="choice ghost" data-choice="The call is up">The call is up</button>
<p class="quiz-note" hidden>RFC 3261. The 200 accepts the session.</p></div>
<div class="quiz" data-quiz data-answer="DSCP 46"><p>What is EF?</p>
<button type="button" class="choice ghost" data-choice="DSCP 46">DSCP 46</button>
<button type="button" class="choice ghost" data-choice="A CUCM partition">A CUCM partition</button>
<p class="quiz-note" hidden>RFC 3246. Telephony uses this mark in RFC 4594.</p></div>
<div class="quiz" data-quiz data-answer="Two dialogs"><p>What does a B2BUA create?</p>
<button type="button" class="choice ghost" data-choice="Two dialogs">Two dialogs</button>
<button type="button" class="choice ghost" data-choice="One proxy transaction">One proxy transaction</button>
<p class="quiz-note" hidden>CUBE and CUCM on a trunk are B2BUAs. Count the Call-IDs.</p></div>
<div class="room-downloads"><a class="dl-btn" href="downloads/flashcards.html">Flashcards</a></div>
{play_block}
"""
    )
    write_page("play.html", "Play", "A short UC drill.", "play", body)


def emit_listen() -> None:
    listen_sources = entries(["flagship-audio"])
    listen_block = html_block(listen_sources, "sources/listen/sources.json")
    mirror_sources(ROOT / "sources" / "listen" / "sources.json", "listen", listen_sources)
    write_json("listen.json", "listen", listen_sources)
    body = page_wrap(
        f"""
<h1>Listen</h1>
<p>The phone bar and the guide can read this page aloud. Recorded desk notes from the flagship studio are optional and separate from the guide voice.</p>
<audio controls preload="none" src="flagship/audio/architect-tour.mp3"><a href="flagship/audio/architect-tour.mp3">Tour notes</a></audio>
<audio controls preload="none" src="flagship/audio/architect-practice.mp3"><a href="flagship/audio/architect-practice.mp3">Practice notes</a></audio>
<p><a href="flagship/audio/scripts/architect-tour.md">Tour script</a></p>
<div class="room-downloads"><button type="button" class="dl-btn" data-listen>Read this page</button></div>
{listen_block}
"""
    )
    write_page("listen.html", "Listen", "Read-along and recorded notes.", "listen", body)


def emit_ecosystem(room_rows: list[dict]) -> None:
    terms = []
    seen = set()
    for room in room_rows:
        for term, meaning in room["glossary"]:
            if term in seen:
                continue
            seen.add(term)
            terms.append(f"<li><strong>{esc(term)}</strong> — {esc(meaning)}</li>")
    sisters = [
        ("lab-sip", "SIP campus. The page title that opened is SIP Universe Free University."),
        ("lab-webex", "Webex campus. The page title that opened is Webex Lab Free University."),
        ("lab-microsoft", "Microsoft voice campus. The page title that opened is Microsoft Voice and Contact Center Lab."),
        ("lab-connect", "Amazon Connect campus. The page title that opened is Amazon Connect Lab."),
        ("lab-genesys", "Genesys campus. The page title that opened is Genesys Cloud CX Lab."),
        ("lab-five9", "Five9 campus. The page title that opened is Five9 Lab."),
        ("lab-nice", "NICE campus. The page title that opened is Contact Path · NICE CXone Lab."),
        ("lab-ai", "AI campus. The page title that opened is AI omni-channel learning lab."),
        ("lab-uc", "The published copy of this UC course. The page title that opened is UC Lab Free University."),
    ]
    sister_rows = []
    for key, note in sisters:
        item = CATALOG[key]
        sister_rows.append(
            f'<li><a href="{esc(item["url"])}">{esc(item["title"])}</a> — {esc(note)}</li>'
        )
    source_items = entries(
        union(
            ["rfc3261", "rfc3550", "rfc4566", "fcc-mlts", "cer", "cer-plan", "cucm-rec", "rfc7866", "rfc7245", "rfc8825", "rfc7118", "wiki-siptrunk", "wiki-webrtc", "fs-manual"],
            [key for key, _note in sisters],
            ["kari-statute", "local-emergency-law", "recording-rule", "cpaas-matrix"],
        )
    )
    mirror_sources(ROOT / "sources" / "ecosystem" / "sources.json", "ecosystem", source_items)
    write_json("ecosystem.json", "ecosystem", source_items)
    deeper = html_block(source_items, "sources/ecosystem/sources.json", md_href="downloads/ecosystem.md")
    write_text(
        ROOT / "downloads" / "ecosystem.md",
        "# Ecosystem\n\nGlossary, mind map, sister labs, integration map, and future hooks.\n\n"
        + markdown_block(source_items, "../sources/ecosystem/sources.json"),
    )
    body = page_wrap(
        f"""
<h1>Ecosystem</h1>
<p class="lede">Words, drawings, and the sibling labs this course can point at. Counts and titles below are from pages opened in 2026-10 or from this repository. A confirm row has no link.</p>
<section><h2>Glossary</h2><p>{len(terms)} terms from the rooms in this repository. The same term is listed once.</p><ul>{''.join(terms)}</ul></section>
<section><h2>Mind map</h2><div class="diagram-frame"><img src="assets/diagrams/university/glossary-map.svg" alt="Glossary mind map"></div>
<p>Dialog, ELIN, fork, and WebRTC are the four words on the drawing. The list above is the rest.</p></section>
<section><h2>Sister technologies</h2>
<p>SIP, RTP, a CUCM dial plan, IOS dial peers, Expressway traversal, DiffServ, UCCE scripts, Emergency Responder locations, recording forks, and an open PBX dialplan are sisters. Learn the object each word names.</p>
<p>These sibling labs opened in 2026-10. They are separate courses. A shared author is not a shared dial plan, and a name here is not an endorsement.</p>
<ul>{''.join(sister_rows)}</ul>
</section>
<section><h2>Integration map</h2><div class="diagram-frame"><img src="assets/diagrams/university/integration-map.svg" alt="Integration map"></div>
<p>CUCM meets Emergency Responder, a recorder, and another PBX on explicit trunks, route points, and labels. UCCE still meets CUCM on a label. Webex Calling still meets a local gateway on a SIP trunk. The <a href="tracks/T9.html">emergency</a>, <a href="tracks/T10.html">recording</a>, and <a href="tracks/T11.html">open bridge</a> tracks are the rooms for those three hops.</p></section>
<section><h2>Future hooks</h2><div class="diagram-frame"><img src="assets/diagrams/university/future-hooks.svg" alt="Future hooks"></div>
<p>Releases, carriers, and any emergency or recording rule this lab did not open stay outside the course until you confirm them. The AI lab is a sibling for agents and tools, not a UC feature. The SIP, Webex, Microsoft, and contact-center labs are the next rooms when the peer is no longer CUCM.</p></section>
<div class="room-downloads"><a class="dl-btn" href="assets/diagrams/university/glossary-map.svg" download>SVG</a><a class="dl-btn" href="downloads/ecosystem.md">Markdown</a></div>
{deeper}
"""
    )
    write_page("ecosystem.html", "Ecosystem", "Glossary, maps, and sibling labs for the UC course.", "ecosystem", body)


def emit_help_legal() -> None:
    help_sources = entries(["rfc3261", "fcc-mlts", "cucm-rec", "rfc7245"])
    help_block = html_block(help_sources, "sources/help/sources.json")
    mirror_sources(ROOT / "sources" / "help" / "sources.json", "help", help_sources)
    write_json("help.json", "help", help_sources)
    legal_sources = entries(["fcc-mlts", "cer", "cucm-rec", "rfc7245", "kari-statute", "local-emergency-law", "recording-rule"])
    legal_block = html_block(legal_sources, "sources/legal/sources.json")
    mirror_sources(ROOT / "sources" / "legal" / "sources.json", "legal", legal_sources)
    write_json("legal.json", "legal", legal_sources)
    help_body = page_wrap(
        f"""
<h1>Help</h1>
<section><h2>Menu</h2><p>The Menu button is at the top left. The drawer slides in. Escape or a tap on the dimmed page closes it. On a phone the same Menu is the left side of the bottom bar.</p></section>
<section><h2>Course</h2><p>Learn opens a ten-minute SIP path and the full course. Each room has a five-minute block, a drawing, a worked example, a quiz, and sources. Emergency calling and recording are study notes, not legal advice.</p></section>
<section><h2>Bot</h2><p>The monkey greets once after the tour, then sits at the right. Drag it. It can read the page aloud, take a typed question, or try the mic. Ambient sound is off until you turn it on. After a few quiet minutes it may offer help. Not now dismisses that offer. Stop offering keeps it quiet.</p>
<p><button type="button" data-tour-open>Replay the tour</button></p></section>
<section><h2>Downloads</h2><p>Every room has SVG, PNG, HTML, and a cheat sheet. The Downloads page lists the deck, the cards, the templates, and one zip.</p></section>
<div class="room-downloads"><a class="dl-btn" href="downloads.html">Open downloads</a></div>
{help_block}
"""
    )
    legal_body = page_wrap(
        f"""
<h1>Legal and credits</h1>
<section><h2>Copyright</h2><p>The words and drawings in this round-2 course are a student project, © 2026 CypherMonkeyLabs. The deep campus and flagship keep their own notices.</p></section>
<section><h2>Trademarks</h2><p>Cisco, CUCM, CUBE, SRST, Expressway, Emergency Responder, Unified Contact Center Enterprise, Webex, and Webex Calling are trademarks of their owners. Asterisk, FreeSWITCH, and other names on the bridge pages belong to their owners. This lab is not affiliated with or endorsed by them.</p></section>
<section><h2>Not legal advice</h2><p>Emergency-calling and call-recording chapters explain documents this lab opened. They are not legal advice, not an emergency plan, and not a retention policy. Confirm the rule where the system operates.</p></section>
<section><h2>Licences</h2><p>Code written for this course is MIT, matching the repository licence. Vendored libraries under flagship/assets/vendor keep MIT, ISC, and BSD-style notices beside the files. Drawings in this course are shared CC BY-NC-SA unless a file says otherwise. RFC text stays with the IETF and the RFC Editor.</p></section>
<section><h2>Privacy</h2><p>There is no account and no tracker. Progress, the tour flag, voice settings, and badge names stay in this browser.</p></section>
<section><h2>Accessibility</h2><p>Display sets font size and font colour. Bright is the default. Dark is optional. The tour and the drawer work from the keyboard. Escape closes them.</p></section>
<section><h2>Removal</h2><p>Write via <a href="{CONTACT}">{CONTACT}</a> and the page will be corrected or removed.</p></section>
{legal_block}
"""
    )
    write_page("help.html", "Help", "How to use the UC lab.", "help", help_body)
    write_page("legal.html", "Legal and credits", "Copyright, trademarks, privacy, and removal.", "help", legal_body)


def emit_quizzes(room_rows: list[dict]) -> None:
    blocks = []
    for room in room_rows:
        for prompt, why, choices, answer in room["quiz"]:
            buttons = "".join(
                f'<button type="button" class="choice ghost" data-choice="{esc(choice)}">{esc(choice)}</button>' for choice in choices
            )
            blocks.append(
                f'<div class="quiz" data-quiz data-answer="{esc(answer)}"><p><a href="rooms/{room["slug"]}.html">{esc(room["title"])}</a>. {esc(prompt)}</p>{buttons}<p class="quiz-note" hidden>{esc(why)}</p></div>'
            )
    quiz_sources = entries(union(*[ROOM_KEYS[room["slug"]] for room in room_rows]))
    quiz_block = html_block(quiz_sources, "sources/quizzes/sources.json")
    mirror_sources(ROOT / "sources" / "quizzes" / "sources.json", "quizzes", quiz_sources)
    write_json("quizzes.json", "quizzes", quiz_sources)
    write_page("quizzes.html", "Quizzes", "Room quizzes in one list.", "play", page_wrap(f"<h1>Quizzes</h1>{''.join(blocks)}{quiz_block}"))


def cards_for(room_rows: list[dict]) -> list[dict]:
    cards = []
    for room in room_rows:
        for term, meaning in room["glossary"]:
            cards.append({"front": term, "back": meaning, "source": room["sources"][0][0], "room": room["slug"]})
        cards.append({"front": room["title"], "back": room["five"][0], "source": room["sources"][0][0], "room": room["slug"]})
        for prompt, why, _choices, answer in room["quiz"]:
            cards.append({"front": prompt, "back": f"{answer}. {why}", "source": room["sources"][0][0], "room": room["slug"]})
    return cards


def card_sources(card: dict) -> list[dict]:
    slug = card.get("room") or ""
    if slug in ROOM_KEYS:
        return entries(ROOM_KEYS[slug])
    return entries(CORE)


def emit_flashcards(cards: list[dict]) -> None:
    folder = ROOT / "downloads"
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "flashcards.json").write_text(json.dumps({"cards": cards}, indent=2), encoding="utf-8")
    deck_items = entries(CORE)
    write_json("flashcards.json", "flashcards", deck_items)
    mirror_sources(folder / "sources" / "flashcards" / "sources.json", "flashcards", deck_items)
    rows = []
    md_parts = ["# UC flashcards", ""]
    for card in cards:
        room_slug = card.get("room") or "flashcards"
        card_href = f"../rooms/sources/{room_slug}/sources.json" if room_slug != "flashcards" else "sources/flashcards/sources.json"
        block = html_block(card_sources(card), card_href)
        rows.append(
            f"<article><h2>{esc(card['front'])}</h2><p>{esc(card['back'])}</p><p>Source: {esc(card['source'])}</p>{block}</article>"
        )
        md_parts.append(f"## {card['front']}")
        md_parts.append("")
        md_parts.append(card["back"])
        md_parts.append("")
        md_parts.append(markdown_block(card_sources(card), card_href))
    closing = html_block(deck_items, "sources/flashcards/sources.json", md_href="flashcards.md")
    doc = (
        "<!DOCTYPE html><html lang=\"en\"><head><meta charset=\"utf-8\"><title>UC flashcards</title>"
        "<link rel=\"stylesheet\" href=\"../assets/css/footer.css\">"
        "<style>body{font-family:system-ui,sans-serif;margin:24px;background:#fff;color:#0f172a}article{break-inside:avoid;border:1px solid #99f6e4;border-radius:12px;padding:12px;margin:8px 0}</style>"
        f"</head><body><h1>UC flashcards</h1>{''.join(rows)}{closing}{footer()}</body></html>"
    )
    (folder / "flashcards.html").write_text(doc, encoding="utf-8")
    (folder / "flashcards-print.html").write_text(doc, encoding="utf-8")
    md_parts.append(markdown_block(deck_items, "sources/flashcards/sources.json"))
    write_text(folder / "flashcards.md", "\n".join(md_parts))
    preview = []
    for card in cards[:12]:
        room_slug = card.get("room") or "flashcards"
        card_href = f"rooms/sources/{room_slug}/sources.json" if room_slug != "flashcards" else "downloads/sources/flashcards/sources.json"
        block = html_block(card_sources(card), card_href)
        preview.append(f'<article class="room-card"><h2>{esc(card["front"])}</h2><p>{esc(card["back"])}</p>{block}</article>')
    write_page(
        "flashcards.html",
        "Flashcards",
        "Know or again, in this browser.",
        "play",
        page_wrap(
            f"<h1>Flashcards</h1><p>{len(cards)} cards. The deck file is a download. This page shows the first questions from the same list.</p>"
            + "".join(preview)
            + '<div class="room-downloads"><a class="dl-btn" href="downloads/flashcards.json" download>JSON</a><a class="dl-btn" href="downloads/flashcards.html">HTML</a><a class="dl-btn" href="downloads/flashcards.pdf">PDF</a><a class="dl-btn" href="downloads/flashcards.md">Markdown</a></div>'
            + html_block(deck_items, "downloads/sources/flashcards/sources.json", md_href="downloads/flashcards.md")
        ),
    )


def emit_office(cards: list[dict]) -> None:
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    templates = ROOT / "downloads" / "templates"
    templates.mkdir(parents=True, exist_ok=True)
    sheets = ROOT / "downloads" / "sheets"

    room_rows = rooms()

    def wrap_line(line: str, width: int = 90) -> list[str]:
        text = " ".join(line.split())
        if not text:
            return []
        out: list[str] = []
        while text:
            if len(text) <= width:
                out.append(text)
                break
            cut = text.rfind(" ", 0, width)
            if cut < 24:
                cut = width
            out.append(text[:cut].strip())
            text = text[cut:].strip()
        return out

    def pdf(path: Path, title: str, lines: list[str], source_items: list[dict] | None = None) -> None:
        c = canvas.Canvas(str(path), pagesize=letter)
        _width, height = letter
        y = height - 54

        def draw(text: str, size: int = 11, bold: bool = False) -> None:
            nonlocal y
            c.setFont("Helvetica-Bold" if bold else "Helvetica", size)
            for chunk in wrap_line(text, 90 if size >= 11 else 100):
                if y < 54:
                    c.showPage()
                    c.setFont("Helvetica-Bold" if bold else "Helvetica", size)
                    y = height - 54
                c.drawString(48, y, chunk)
                y -= size + 4

        draw(title[:120], 16, True)
        y -= 8
        for line in lines:
            draw(line)
        y -= 8
        draw("Independent free learning-only student project. Trademarks belong to their owners. Not affiliated or endorsed. Contact: https://cyphermonkeylabs.com/#contact. Nothing is sold.", 9)
        if source_items:
            y -= 8
            for line in plain_lines(source_items):
                draw(line, 9)
        c.save()

    for slug, (title, lines) in SHEETS.items():
        source_items = entries(keys_for_sheet(room_rows, slug))
        pdf(sheets / f"{slug}.pdf", title, lines, source_items)
    flash_lines = []
    for card in cards:
        flash_lines.append(f"{card['front']} — {card['back']}")
        flash_lines.extend(plain_lines(card_sources(card)))
    pdf(ROOT / "downloads" / "flashcards.pdf", "UC flashcards", flash_lines, entries(CORE))

    template_text = {
        "evidence-pack": ["Evidence pack", "Hypothesis:", "Time:", "Calling number:", "Called number:", "One change:", "Result:", "Revert plan:"],
        "cutover-checklist": ["Cutover checklist", "Number moved:", "Old route saved:", "Test call passed:", "Caller identity checked:", "Rollback owner:", "Emergency test window:"],
        "dial-plan-worksheet": ["Dial plan worksheet", "Partition:", "CSS order:", "Pattern:", "Route list:", "Gateway or trunk:", "Called transform:", "Calling transform:"],
    }
    for slug, lines in template_text.items():
        source_items = entries(keys_for_sheet(room_rows, TEMPLATE_SHEET[slug]))
        write_json(f"templates/{slug}.json", slug, source_items)
        mirror_sources(templates / "sources" / slug / "sources.json", slug, source_items)
        deeper = html_block(source_items, f"sources/{slug}/sources.json", md_href=f"{slug}.md")
        html_doc = (
            "<!DOCTYPE html><html lang=\"en\"><head><meta charset=\"utf-8\">"
            f"<title>{esc(lines[0])}</title><link rel=\"stylesheet\" href=\"../../assets/css/footer.css\"></head><body>"
            f"<h1>{esc(lines[0])}</h1><ul>" + "".join(f"<li>{esc(line)}</li>" for line in lines[1:]) + f"</ul>{deeper}{footer()}</body></html>"
        )
        (templates / f"{slug}.html").write_text(html_doc, encoding="utf-8")
        pdf(templates / f"{slug}.pdf", lines[0], lines[1:], source_items)
        document = Document()
        document.add_heading(lines[0], 0)
        for line in lines[1:]:
            document.add_paragraph(line)
        document.add_paragraph("Independent, free, learning-only student project. Trademarks belong to their owners. Not affiliated or endorsed. Concerns via https://cyphermonkeylabs.com/#contact. Nothing is sold.")
        add_docx_sources(document, source_items)
        document.save(templates / f"{slug}.docx")
        write_text(
            templates / f"{slug}.md",
            sheet_markdown(lines[0], lines[1:], source_items, f"sources/{slug}/sources.json"),
        )

    pres = Presentation()
    slides = [
        ("UC Lab Free University", "SIP, CUCM, CUBE, SRST, Expressway, QoS, UCCE, Webex Calling migration."),
        ("SIP", "A transaction is not a dialog. 180 alerts. 200 accepts. ACK finishes."),
        ("Dial plan", "Partitions, ordered CSS, route list, route group."),
        ("CUBE", "Two dialogs. Read both legs."),
        ("SRST", "A planned subset when CUCM is unreachable."),
        ("Expressway", "C, E, and a traversal zone. Confirm ports."),
        ("QoS", "EF is DSCP 46. LLQ does not create bandwidth."),
        ("UCCE", "Script, label, then a number CUCM can route."),
        ("Migration", "Dual-run a number you can return."),
        ("Troubleshooting", "One hypothesis. One change. Save the ladder."),
        ("Emergency calling", "Study notes. Direct dial, Emergency Responder, ELIN callback. Not legal advice."),
        ("Call recording", "Phone fork or gateway fork. SIPREC is a separate session."),
        ("Open bridges", "SIP trunk, Asterisk or FreeSWITCH, WebRTC through a gateway."),
        ("Disclaimer", "Independent free learning project. Not affiliated or endorsed. Not legal advice. Contact cyphermonkeylabs.com. Nothing is sold."),
    ]
    for title, line in slides:
        layout = pres.slide_layouts[1]
        slide = pres.slides.add_slide(layout)
        slide.shapes.title.text = title
        slide.placeholders[1].text = line
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(20)
                    run.font.name = "Calibri"
    add_pptx_sources(pres, entries(CORE))
    pres.save(ROOT / "downloads" / "uc-lab-deck.pptx")


def emit_catalog(room_rows: list[dict]) -> None:
    pages = [
        {"title": "Home", "href": "index.html", "blurb": f"{len(TRACKS)} tracks", "keys": ["home", "start"]},
        {"title": "Learn", "href": "start.html", "blurb": "Ten minutes or the full course", "keys": ["learn", "course"]},
        {"title": "Diagrams", "href": "diagrams.html", "blurb": "Drawings", "keys": ["diagram"]},
        {"title": "Flows", "href": "flows.html", "blurb": "Ladders", "keys": ["flow", "ladder"]},
        {"title": "Play", "href": "play.html", "blurb": "Drill", "keys": ["play", "quiz"]},
        {"title": "Listen", "href": "listen.html", "blurb": "Read along", "keys": ["listen", "audio"]},
        {"title": "Downloads", "href": "downloads.html", "blurb": "Exports", "keys": ["download", "export"]},
        {"title": "Ecosystem", "href": "ecosystem.html", "blurb": "Glossary", "keys": ["glossary", "ecosystem"]},
        {"title": "Help", "href": "help.html", "blurb": "Tour and bot", "keys": ["help", "tour"]},
        {"title": "Legal", "href": "legal.html", "blurb": "Credits", "keys": ["legal", "privacy"]},
    ]
    for room in room_rows:
        pages.append(
            {
                "title": room["title"],
                "href": f"rooms/{room['slug']}.html",
                "blurb": room["five"][0],
                "keys": [room["slug"].replace("-", " "), room["id"].lower(), track_key(room["track"])],
            }
        )
    path = ROOT / "assets" / "data"
    path.mkdir(parents=True, exist_ok=True)
    (path / "catalog.json").write_text(json.dumps({"pages": pages}, indent=2), encoding="utf-8")


def track_key(tid: str) -> str:
    for row in TRACKS:
        if row[0] == tid:
            return row[1].lower()
    return tid


def emit_downloads(items: list[dict]) -> None:
    rows = []
    interesting = []
    for folder in (ROOT / "downloads", ROOT / "assets" / "diagrams" / "university"):
        for path in sorted(folder.rglob("*")):
            if not path.is_file():
                continue
            if path.suffix.lower() not in {".html", ".pdf", ".png", ".svg", ".json", ".docx", ".pptx", ".zip", ".md"}:
                continue
            interesting.append(path)
    # The zip is added after this table is first built. Include planned zip if present.
    for path in interesting:
        text = ""
        if path.suffix.lower() in {".html", ".svg", ".json"}:
            text = path.read_text(encoding="utf-8", errors="ignore")
        has_note = "learning-only" in text or "Sources and go deeper" in text or path.suffix.lower() in {".pdf", ".docx", ".pptx", ".png", ".zip"}
        rel = path.relative_to(ROOT).as_posix()
        rows.append(f"<tr><td>{esc(rel)}</td><td>{path.stat().st_size}</td><td>{'yes' if has_note else 'check'}</td></tr>")
    body = page_wrap(
        f"""
<h1>Downloads</h1>
<p><a class="btn" href="index.html">Open app</a> <a class="ghost" href="{ORG}">View repo</a></p>
<div class="room-downloads">
<a class="dl-btn" href="downloads/uc-lab-deck.pptx" download>PPTX deck</a>
<a class="dl-btn" href="downloads/flashcards.json" download>Flashcards JSON</a>
<a class="dl-btn" href="downloads/flashcards.pdf" download>Flashcards PDF</a>
<a class="dl-btn" href="downloads/flashcards.html">Flashcards HTML</a>
<a class="dl-btn" href="downloads/templates/evidence-pack.docx" download>Evidence DOCX</a>
<a class="dl-btn" href="downloads/uc-lab-university.zip" download>Everything zip</a>
</div>
<h2>Cheat sheets</h2>
<div class="room-downloads">
{''.join(f'<a class="dl-btn" href="downloads/sheets/{esc(slug)}.html">{esc(slug)} HTML</a><a class="dl-btn" href="downloads/sheets/{esc(slug)}.pdf">{esc(slug)} PDF</a><a class="dl-btn" href="downloads/sheets/{esc(slug)}.md">{esc(slug)} MD</a><a class="dl-btn" href="downloads/sheets/{esc(slug)}.png">{esc(slug)} PNG</a>' for slug in SHEETS)}
</div>
<h2>Templates</h2>
<div class="room-downloads">
<a class="dl-btn" href="downloads/templates/evidence-pack.html">Evidence HTML</a>
<a class="dl-btn" href="downloads/templates/cutover-checklist.pdf">Cutover PDF</a>
<a class="dl-btn" href="downloads/templates/dial-plan-worksheet.docx">Dial plan DOCX</a>
</div>
<h2>Headless test table</h2>
<p>Sizes are bytes on disk after the build. HTML and SVG rows contain the learning-only disclaimer. PDF, DOCX, and PPTX include the same note in the document text. PNG files are raster copies of those pages.</p>
<div class="table-wrap"><table class="plain"><thead><tr><th>File</th><th>Bytes</th><th>Disclaimer</th></tr></thead><tbody>
{''.join(rows)}
</tbody></table></div>
<p>{len(items)} diagrams are linked from <a href="diagrams.html">Diagrams</a>.</p>
{html_block(entries(CORE), "sources/downloads/sources.json")}
"""
    )
    mirror_sources(ROOT / "sources" / "downloads" / "sources.json", "downloads", entries(CORE))
    write_json("downloads.json", "downloads", entries(CORE))
    write_page("downloads.html", "Downloads", "Sheets, cards, templates, and the zip.", "downloads", body)


def render_pngs() -> None:
    import subprocess
    from concurrent.futures import ThreadPoolExecutor

    pairs = []
    for folder in (ROOT / "downloads" / "diagrams", ROOT / "downloads" / "sheets"):
        for html_path in sorted(folder.glob("*.html")):
            pairs.append((html_path, html_path.with_suffix(".png")))

    def shot(pair: tuple[Path, Path]) -> None:
        html_path, png_path = pair
        profile = f"/tmp/ucpng-{png_path.stem}"
        subprocess.run(
            [
                "/usr/bin/google-chrome-stable",
                "--headless=new",
                "--no-sandbox",
                "--disable-gpu",
                f"--user-data-dir={profile}",
                "--hide-scrollbars",
                "--window-size=1100,800",
                f"--screenshot={png_path}",
                html_path.as_uri(),
            ],
            check=True,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )

    with ThreadPoolExecutor(max_workers=4) as pool:
        list(pool.map(shot, pairs))


def emit_zip() -> None:
    target = ROOT / "downloads" / "uc-lab-university.zip"
    include_roots = ["index.html", "start.html", "diagrams.html", "flows.html", "downloads.html", "help.html", "legal.html", "assets", "guide", "rooms", "tracks", "downloads", "sw.js", "manifest.webmanifest", "favicon.svg"]
    with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name in include_roots:
            path = ROOT / name
            if path.is_file():
                archive.write(path, name)
            elif path.is_dir():
                for file in path.rglob("*"):
                    if not file.is_file():
                        continue
                    if "uc-lab-university.zip" in file.name:
                        continue
                    if file.suffix == ".png" and "screenshots" in file.parts:
                        continue
                    archive.write(file, file.relative_to(ROOT).as_posix())
    page = ROOT / "downloads.html"
    if page.exists():
        size = target.stat().st_size
        text = page.read_text(encoding="utf-8")
        updated, count = re.subn(
            r"(uc-lab-university\.zip</td><td>)\d+",
            rf"\g<1>{size}",
            text,
            count=1,
        )
        if count and updated != text:
            page.write_text(updated, encoding="utf-8")


def emit_source_extras() -> None:
    """Flagship cheat HTML, older flagship PDFs, and the campus lesson export."""
    room_rows = rooms()
    patch_flagship_cheats()
    flagship_pdfs = {
        "cube-sbc.pdf": keys_for_sheet(room_rows, "cube"),
        "cucm-dial-plan.pdf": keys_for_sheet(room_rows, "dial-plan"),
        "cucm-webex-migration.pdf": keys_for_sheet(room_rows, "migration"),
        "e911-design.pdf": keys_for_sheet(room_rows, "e911"),
        "qos-dscp.pdf": keys_for_sheet(room_rows, "qos"),
        "sdp-offer-answer.pdf": ROOM_KEYS["offer-answer"],
        "sip-call-flows.pdf": keys_for_sheet(room_rows, "sip"),
    }
    for name, keys in flagship_pdfs.items():
        path = ROOT / "flagship" / "downloads" / name
        if path.is_file():
            append_pdf_sources(path, entries(keys))
    cheat_dir = ROOT / "flagship" / "cheats"
    for slug, keys in CHEAT_KEYS.items():
        pdf_path = cheat_dir / f"{slug}.pdf"
        if pdf_path.is_file():
            append_pdf_sources(pdf_path, entries(keys))
    write_campus_bundle(entries(CORE))


def scrub_legacy() -> int:
    changed = 0
    email = re.compile(r"([A-Za-z0-9._%+-]+)@([A-Za-z0-9-]+\.[A-Za-z]{2,})")
    for path in ROOT.rglob("*"):
        if not path.is_file():
            continue
        if path.suffix.lower() not in {".html", ".md", ".js", ".css", ".json"}:
            continue
        if any(part in {".git", "vendor", "node_modules"} for part in path.parts):
            continue
        if path.relative_to(ROOT).as_posix() == "campus/index.html":
            continue
        text = path.read_text(encoding="utf-8", errors="ignore")
        original = text
        if path.suffix.lower() == ".html" and "data-uc-page" not in text:
            text = re.sub(r'<nav class="brand-links".*?</nav>', "", text, flags=re.S)
            text = email.sub(r"\1 at \2", text)
        text = text.replace("https://github.com/cipher0x9", ORG)
        text = re.sub(r"https://x\.com/[^\s\"'<)]+", WEB, text)
        text = re.sub(r"https://huggingface\.co/[^\s\"'<)]+", WEB, text)
        text = re.sub(r"https://(www\.)?linkedin\.com/[^\s\"'<)]+", WEB, text)
        text = re.sub(r"mailto:[^\s\"'<)]+", CONTACT, text)
        if path.suffix.lower() == ".html" and "data-uc-page" not in text and "learning-only student project" not in text:
            depth = len(path.relative_to(ROOT).parts) - 1
            pre = "../" * depth
            snippet = (
                f'<p class="uc-legacy-dl"><a href="{pre}downloads.html">Download exports</a></p>'
                + footer()
            )
            if "</body>" in text:
                text = text.replace("</body>", snippet + "</body>", 1)
            else:
                text += snippet
        if text != original:
            path.write_text(text, encoding="utf-8")
            changed += 1
    return changed


def main() -> None:
    import sys
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    items = build_diagrams()
    room_rows = rooms()
    emit_diagrams(items)
    emit_sheets()
    emit_rooms(room_rows)
    emit_tracks(room_rows)
    emit_home(room_rows, items)
    emit_start(room_rows)
    emit_diagram_indexes(items)
    emit_charts(room_rows, items)
    emit_play()
    emit_listen()
    emit_ecosystem(room_rows)
    emit_help_legal()
    emit_quizzes(room_rows)
    card_rows = cards_for(room_rows)
    emit_flashcards(card_rows)
    emit_office(card_rows)
    emit_catalog(room_rows)
    render_pngs()
    emit_downloads(items)
    scrubbed = scrub_legacy()
    emit_source_extras()
    emit_zip()
    print(f"rooms {len(room_rows)} diagrams {len(items)} cards {len(card_rows)} scrubbed {scrubbed}")


if __name__ == "__main__":
    main()
