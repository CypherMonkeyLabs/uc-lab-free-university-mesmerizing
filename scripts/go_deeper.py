#!/usr/bin/env python3
"""Verified 'Sources and go deeper' blocks for rooms and downloadables.

Every URL in CATALOG was opened in October 2026. A confirm row has no URL.
Do not add a link that was not opened.
"""
from __future__ import annotations

import html
import json
import re
from pathlib import Path

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx.util import Pt
from pypdf import PdfReader, PdfWriter
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent.parent
CHECKED = "2026-10"
DATA = ROOT / "assets" / "data" / "sources"

KIND_LABEL = {
    "rfc": "RFC",
    "cisco": "Cisco",
    "ports": "Protocol and ports",
    "community": "Community",
    "wikipedia": "Wikipedia",
    "itu": "ITU",
    "webex": "Webex",
}


def _row(kind: str, title: str, url: str | None, note: str = "") -> dict:
    return {
        "kind": kind,
        "title": title,
        "url": url,
        "status": "verified" if url else "confirm",
        "last_checked": CHECKED,
        "note": note,
    }


CATALOG: dict[str, dict] = {
    "rfc3261": _row("rfc", "RFC 3261, SIP", "https://www.rfc-editor.org/rfc/rfc3261"),
    "rfc3262": _row("rfc", "RFC 3262, reliable provisional responses (PRACK)", "https://www.rfc-editor.org/rfc/rfc3262"),
    "rfc3264": _row("rfc", "RFC 3264, offer/answer", "https://www.rfc-editor.org/rfc/rfc3264"),
    "rfc3515": _row("rfc", "RFC 3515, SIP REFER", "https://www.rfc-editor.org/rfc/rfc3515"),
    "rfc3550": _row("rfc", "RFC 3550, RTP", "https://www.rfc-editor.org/rfc/rfc3550"),
    "rfc3711": _row("rfc", "RFC 3711, SRTP", "https://www.rfc-editor.org/rfc/rfc3711"),
    "rfc4028": _row("rfc", "RFC 4028, SIP session timers", "https://www.rfc-editor.org/rfc/rfc4028"),
    "rfc4566": _row("rfc", "RFC 4566, SDP", "https://www.rfc-editor.org/rfc/rfc4566"),
    "rfc4733": _row("rfc", "RFC 4733, RTP telephony events", "https://www.rfc-editor.org/rfc/rfc4733"),
    "rfc3246": _row("rfc", "RFC 3246, Expedited Forwarding", "https://www.rfc-editor.org/rfc/rfc3246"),
    "rfc4594": _row("rfc", "RFC 4594, DiffServ service classes", "https://www.rfc-editor.org/rfc/rfc4594"),
    "rfc7866": _row("rfc", "RFC 7866, session recording protocol", "https://www.rfc-editor.org/rfc/rfc7866"),
    "rfc8445": _row("rfc", "RFC 8445, ICE", "https://www.rfc-editor.org/rfc/rfc8445"),
    "cucm": _row("cisco", "Cisco Unified Communications Manager product support", "https://www.cisco.com/c/en/us/support/unified-communications/unified-communications-manager-callmanager/series.html"),
    "cube": _row("cisco", "Cisco Unified Border Element product support", "https://www.cisco.com/c/en/us/support/unified-communications/unified-border-element/series.html"),
    "srst": _row("cisco", "Cisco Unified SRST product support", "https://www.cisco.com/c/en/us/support/unified-communications/unified-survivable-remote-site-telephony/series.html"),
    "exp": _row("cisco", "Cisco Expressway Series product support", "https://www.cisco.com/c/en/us/support/unified-communications/expressway-series/series.html"),
    "ucce": _row("cisco", "Cisco Unified Contact Center Enterprise product support", "https://www.cisco.com/c/en/us/support/customer-collaboration/unified-contact-center-enterprise/series.html"),
    "cucm-sys": _row("cisco", "System Configuration Guide for CUCM, Release 15", "https://www.cisco.com/c/en/us/td/docs/voice_ip_comm/cucm/admin/15/systemConfig/cucm_b_system-configuration-guide-15.html"),
    "pa-index": _row("cisco", "Cisco Collaboration Preferred Architectures", "https://www.cisco.com/c/en/us/td/docs/voice_ip_comm/uc_system/design/guides/PAdocs.html"),
    "pa15": _row("cisco", "Preferred Architecture for Cisco Collaboration 15, design overview (PDF)", "https://www.cisco.com/c/dam/en/us/td/docs/solutions/PA/overview/15x/clbpa15x.pdf"),
    "dialpeer": _row("cisco", "Understand IOS and IOS XE call routing", "https://www.cisco.com/c/en/us/support/docs/voice/ip-telephony-voice-over-ip-voip/211306-In-Depth-Explanation-of-Cisco-IOS-and-IO.html"),
    "cube-dp": _row("cisco", "CUBE configuration guide, dial peers (IOS XE 17.6 onwards)", "https://www.cisco.com/c/en/us/td/docs/ios-xml/ios/voice/cube/ios-xe/config/ios-xe-book/m_cube-dp.html"),
    "mra": _row("cisco", "Mobile and Remote Access through Cisco Expressway, X15.2", "https://www.cisco.com/c/en/us/td/docs/voice_ip_comm/expressway/config_guide/X15-2/mra/exwy_b_mra-deployment-guide-x152.html"),
    "cucm-ports": _row("ports", "CUCM 15 TCP and UDP port usage", "https://www.cisco.com/c/en/us/td/docs/voice_ip_comm/cucm/admin/15/systemConfig/cucm_b_system-configuration-guide-15/cucm_m_tcp-and-udp-port-usage-12-0.html"),
    "exp-ports": _row("ports", "Cisco Expressway IP port usage (X14.3 and X15.0)", "https://www.cisco.com/c/en/us/td/docs/voice_ip_comm/expressway/config_guide/X15-0/ip-port/exwy_b_cisco-expressway-ip-port-usage-configuration-guide-x15.html"),
    "webex-ports": _row("ports", "Port reference for Webex Calling", "https://help.webex.com/en-us/article/b2exve/Port-Reference-Information-for-Cisco-Webex-Calling"),
    "iana-sip": _row("ports", "IANA SIP parameters", "https://www.iana.org/assignments/sip-parameters/sip-parameters.xhtml"),
    "iana-rtp": _row("ports", "IANA RTP parameters", "https://www.iana.org/assignments/rtp-parameters/rtp-parameters.xhtml"),
    "iana-ports": _row("ports", "IANA service names and port numbers", "https://www.iana.org/assignments/service-names-port-numbers/service-names-port-numbers.xhtml"),
    "webex": _row("webex", "Webex Help Center", "https://help.webex.com/"),
    "webex-calling": _row("webex", "Get started with Webex Calling", "https://help.webex.com/en-us/article/8jda92/Webex-Calling"),
    "lgw": _row("webex", "Get started with Local Gateway", "https://help.webex.com/en-us/article/t9xctu/Get-started-with-Local-Gateway"),
    "lgw-ios": _row("webex", "Configure Local Gateway on Cisco IOS XE for Webex Calling", "https://help.webex.com/en-us/article/jr1i3r/Configure-local-gateway-on-IOS-XE-for-Webex-Calling"),
    "g114": _row("itu", "ITU-T G.114, one-way transmission time", "https://www.itu.int/rec/T-REC-G.114"),
    "g711": _row("itu", "ITU-T G.711, pulse code modulation of voice", "https://www.itu.int/rec/T-REC-G.711"),
    "community": _row("community", "Cisco Community, Collaboration knowledge base", "https://community.cisco.com/t5/collaboration-knowledge-base/tkb-p/4691-docs-collaboration-voice-video"),
    "wiki-sip": _row("wikipedia", "Session Initiation Protocol (Wikipedia)", "https://en.wikipedia.org/wiki/Session_Initiation_Protocol"),
    "wiki-rtp": _row("wikipedia", "Real-time Transport Protocol (Wikipedia)", "https://en.wikipedia.org/wiki/Real-time_Transport_Protocol"),
    "wiki-sdp": _row("wikipedia", "Session Description Protocol (Wikipedia)", "https://en.wikipedia.org/wiki/Session_Description_Protocol"),
    "wiki-diff": _row("wikipedia", "Differentiated services (Wikipedia)", "https://en.wikipedia.org/wiki/Differentiated_services"),
    "wiki-qos": _row("wikipedia", "Quality of service (Wikipedia)", "https://en.wikipedia.org/wiki/Quality_of_service"),
    "wiki-voip": _row("wikipedia", "Voice over IP (Wikipedia)", "https://en.wikipedia.org/wiki/Voice_over_IP"),
    "wiki-sccp": _row("wikipedia", "Skinny Client Control Protocol (Wikipedia)", "https://en.wikipedia.org/wiki/Skinny_Call_Control_Protocol"),
    "wiki-cucm": _row("wikipedia", "CUCM (Wikipedia disambiguation)", "https://en.wikipedia.org/wiki/CUCM"),
    "srnd-book": _row("cisco", "Collaboration 12 SRND book", None, "The historical Collaboration 12 SRND path opened as a product listing, not the design book. Use the Preferred Architecture index, which was opened."),
    "qos-srnd": _row("cisco", "Enterprise QoS SRND book", None, "The Enterprise QoS SRND path returned HTTP 404."),
    "wiki-cucm-article": _row("wikipedia", "English Wikipedia article for Cisco Unified Communications Manager", None, "That title returned HTTP 404. The CUCM disambiguation page is the Wikipedia page that opened."),
    "wiki-cube": _row("wikipedia", "English Wikipedia article for Cisco Unified Border Element", None, "That title returned HTTP 404."),
    "wiki-srst": _row("wikipedia", "English Wikipedia article for Survivable Remote Site Telephony", None, "That title returned HTTP 404."),
    "survivability": _row("webex", "Webex Calling survivability gateway guide", None, "A survivability-gateway article was not opened, so this row has no link. The local-gateway guide that did open is listed above."),
    "e911-law": _row("cisco", "Emergency-calling statute for your country", None, "No emergency-law page was opened for this block. Confirm the rule that applies where the site is."),
}


def union(*groups: list[str]) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for group in groups:
        for key in group:
            if key not in seen:
                seen.add(key)
                out.append(key)
    return out


_SIP = ["rfc3261", "rfc3550", "rfc4566", "cucm", "cucm-sys", "cucm-ports", "iana-sip", "iana-rtp", "community", "wiki-sip", "wiki-rtp", "srnd-book"]
_DIAL = ["cucm", "cucm-sys", "pa-index", "pa15", "cucm-ports", "iana-sip", "community", "wiki-voip", "wiki-cucm", "srnd-book", "wiki-cucm-article"]
_CUBE = ["rfc3261", "rfc3550", "rfc3711", "cube", "cube-dp", "dialpeer", "cucm-ports", "iana-sip", "community", "wiki-sip", "wiki-cube"]
_EXP = ["exp", "mra", "exp-ports", "cucm", "cucm-ports", "pa-index", "community", "wiki-voip", "srnd-book"]
_QOS = ["rfc3246", "rfc4594", "g114", "g711", "pa-index", "cucm-sys", "community", "wiki-diff", "wiki-qos", "qos-srnd"]
_UCCE = ["ucce", "cucm", "cucm-sys", "pa-index", "cucm-ports", "community", "wiki-voip", "srnd-book"]
_MOVE = ["webex-calling", "webex", "lgw", "lgw-ios", "webex-ports", "cucm", "cucm-sys", "community", "wiki-voip", "survivability"]
_FIX = ["rfc3261", "rfc3264", "rfc3550", "cube", "dialpeer", "cucm-sys", "cucm-ports", "community", "wiki-sip", "wiki-rtp"]

ROOM_KEYS: dict[str, list[str]] = {
    "sip-transactions": union(_SIP, ["rfc3264"]),
    "offer-answer": union(["rfc3264", "rfc3261", "rfc4566"], _SIP, ["wiki-sdp"]),
    "sip-registration": union(["rfc3261", "cucm", "cucm-sys", "cucm-ports", "iana-sip", "community", "wiki-sip", "wiki-sccp", "wiki-cucm", "wiki-cucm-article"]),
    "prack-early-media": union(["rfc3262", "rfc3261", "rfc4566", "cucm-sys", "iana-sip", "community", "wiki-sip", "wiki-sdp"]),
    "partitions-css": _DIAL,
    "route-patterns": _DIAL,
    "transformations": _DIAL,
    "sip-trunks": union(["rfc3261", "rfc4028", "cucm", "cucm-sys", "cucm-ports", "iana-sip", "community", "wiki-sip", "srnd-book"]),
    "dial-peers": union(["dialpeer", "cube-dp", "rfc4733", "cube", "iana-sip", "iana-ports", "community", "wiki-sip", "wiki-voip", "wiki-cube"]),
    "cube-b2bua": union(_CUBE, ["rfc4733", "rfc4566"]),
    "srst": union(["srst", "cucm", "cucm-sys", "cucm-ports", "community", "wiki-voip", "wiki-srst"]),
    "local-gateway": union(["lgw", "lgw-ios", "webex-calling", "webex-ports", "rfc3711", "rfc3261", "cube", "community", "wiki-voip"]),
    "expressway-pair": _EXP,
    "mra": union(["mra", "exp", "exp-ports", "cucm-sys", "iana-ports", "community", "wiki-voip"]),
    "firewall-traversal": union(["exp-ports", "rfc3550", "rfc8445", "exp", "iana-ports", "community", "wiki-rtp"]),
    "dscp-trust": _QOS,
    "llq-delay": _QOS,
    "locations-cac": union(["cucm-sys", "pa15", "rfc4594", "cucm", "community", "wiki-qos", "qos-srnd"]),
    "ucce-path": _UCCE,
    "cucm-ucce-bridge": union(["ucce", "cucm-sys", "cucm-ports", "community", "wiki-voip"]),
    "wxcc-and-ucce": union(["ucce", "webex", "webex-calling", "community", "wiki-voip"]),
    "calling-versus-di": union(["webex-calling", "webex", "cucm", "pa-index", "community", "wiki-voip"]),
    "dual-run": union(["webex-calling", "cucm-sys", "lgw", "community", "wiki-voip"]),
    "cutover-survivability": union(_MOVE, ["srst"]),
    "one-change": union(["cucm-sys", "rfc3261", "community", "wiki-sip", "cucm-ports"]),
    "read-a-ladder": union(["rfc3261", "rfc3515", "rfc3264", "rfc3550", "cube", "dialpeer", "community", "wiki-sip"]),
    "one-way-audio": union(["rfc3550", "rfc3264", "rfc3711", "rfc8445", "cube", "cucm-ports", "exp-ports", "community", "wiki-rtp"]),
}

CORE = union(
    ["rfc3261", "rfc3264", "rfc3550", "rfc4566", "rfc3711"],
    ["cucm", "cucm-sys", "pa-index", "pa15", "cube"],
    ["cucm-ports", "exp-ports", "webex-ports", "iana-sip", "iana-rtp", "iana-ports"],
    ["community", "wiki-sip", "wiki-rtp", "wiki-sdp", "wiki-voip"],
    ["srnd-book", "qos-srnd", "wiki-cucm-article"],
)

SHEET_EXTRA: dict[str, list[str]] = {
    "e911": union(["webex-calling", "cucm-sys", "cucm-ports", "community", "wiki-voip", "e911-law"]),
    "bandwidth": union(["g711", "rfc3550", "rfc4594", "community", "wiki-qos", "qos-srnd"]),
    "ports-note": union(["cucm-ports", "exp-ports", "webex-ports", "iana-ports", "iana-sip", "iana-rtp", "community", "wiki-rtp"]),
    "evidence": union(["rfc3261", "cucm-sys", "community", "wiki-sip", "cucm-ports"]),
}

CHEAT_KEYS: dict[str, list[str]] = {
    "call-flows": union(["rfc3261", "rfc3264", "rfc3515", "rfc3550", "rfc4566", "cucm-ports", "iana-sip", "community", "wiki-sip", "srnd-book"]),
    "ports": SHEET_EXTRA["ports-note"],
    "firewall": union(["exp-ports", "cucm-ports", "rfc3550", "rfc8445", "community", "wiki-rtp"]),
    "config": union(["cucm-sys", "cucm", "webex-calling", "community", "wiki-voip", "srnd-book"]),
    "troubleshoot": _FIX,
    "glossary": union(["rfc3261", "cucm-sys", "cucm-ports", "community", "wiki-sip", "wiki-voip"]),
    "pstn": union(["rfc3261", "webex-calling", "lgw", "cube", "community", "wiki-voip"]),
    "e911": SHEET_EXTRA["e911"],
    "qos": _QOS,
    "migration": _MOVE,
    "security": union(["rfc3711", "rfc3550", "rfc7866", "cucm-ports", "community", "wiki-rtp"]),
    "wfm-recording": union(["rfc7866", "rfc3550", "cucm-ports", "community", "wiki-rtp"]),
    "flashcards": CORE,
    "index": CORE,
}

TEMPLATE_SHEET = {
    "evidence-pack": "evidence",
    "cutover-checklist": "migration",
    "dial-plan-worksheet": "dial-plan",
}


def entries(keys: list[str]) -> list[dict]:
    missing = [key for key in keys if key not in CATALOG]
    if missing:
        raise KeyError("unknown source " + ", ".join(missing))
    return [dict(CATALOG[key]) for key in keys]


def keys_for_sheet(room_rows: list[dict], sheet: str) -> list[str]:
    if sheet in SHEET_EXTRA:
        return SHEET_EXTRA[sheet]
    groups = [ROOM_KEYS[room["slug"]] for room in room_rows if room["sheet"] == sheet]
    return union(*groups) or CORE


def keys_for_track(room_rows: list[dict], tid: str) -> list[str]:
    groups = [ROOM_KEYS[room["slug"]] for room in room_rows if room["track"] == tid]
    return union(*groups) or CORE


def keys_for_diagram(room_rows: list[dict], slug: str) -> list[str]:
    groups = [ROOM_KEYS[room["slug"]] for room in room_rows if room["diagram"] == slug]
    return union(*groups) or CORE


def _esc(value: str) -> str:
    return html.escape(value, quote=True)


def _intro() -> str:
    return f"Opened {CHECKED}. A row marked confirm has no link, because that page was not verified."


def html_block(items: list[dict], json_href: str, *, dark: bool = False, md_href: str = "") -> str:
    link = "#99f6e4" if dark else "#134e4a"
    style = "margin:28px 0;padding:16px 18px;border:2px solid #0f766e;border-radius:16px;background:#f0fdfa;color:#042f2e"
    if dark:
        style = "margin:28px 0;padding:16px 18px;border:2px solid #5eead4;border-radius:16px;background:#042f2e;color:#ecfeff"
    rows = []
    for item in items:
        label = KIND_LABEL[item["kind"]]
        checked = f"last checked {item['last_checked']}"
        note = f" {_esc(item['note'])}" if item["note"] else ""
        if item["url"]:
            body = f'<a href="{_esc(item["url"])}" style="color:{link}">{_esc(item["title"])}</a>'
        else:
            body = f'{_esc(item["title"])}. confirm.{note}'
        rows.append(f"<li><strong>{label}.</strong> {body} <span style=\"white-space:nowrap\">{checked}</span></li>")
    extra = ""
    if md_href:
        extra += f' <a href="{_esc(md_href)}">Markdown</a>'
    return (
        f'<section class="go-deeper" style="{style}">'
        "<h2>Sources and go deeper</h2>"
        f"<p>{_esc(_intro())}</p>"
        f"<ul>{''.join(rows)}</ul>"
        f'<p><a href="{_esc(json_href)}">sources.json</a>{extra}</p>'
        "</section>"
    )


def markdown_block(items: list[dict], json_href: str) -> str:
    lines = ["## Sources and go deeper", "", "Independent, free, learning-only student project.", "", _intro(), ""]
    for item in items:
        label = KIND_LABEL[item["kind"]]
        checked = f"last checked {item['last_checked']}"
        if item["url"]:
            lines.append(f"- {label}. [{item['title']}]({item['url']}). {checked}")
        else:
            note = f" {item['note']}" if item["note"] else ""
            lines.append(f"- {label}. {item['title']}. confirm.{note} {checked}")
    lines.append("")
    lines.append(f"Room or sheet file: {json_href}")
    lines.append("")
    return "\n".join(lines)


def plain_lines(items: list[dict]) -> list[str]:
    lines = ["Sources and go deeper", _intro()]
    for item in items:
        label = KIND_LABEL[item["kind"]]
        checked = f"last checked {item['last_checked']}"
        if item["url"]:
            lines.append(f"{label}. {item['title']}. {item['url']}. {checked}")
        else:
            note = f" {item['note']}" if item["note"] else ""
            lines.append(f"{label}. {item['title']}. confirm.{note} {checked}")
    return lines


def payload(owner: str, items: list[dict]) -> dict:
    return {"room": owner, "checked": CHECKED, "sources": items}


def write_json(rel: str, owner: str, items: list[dict]) -> None:
    path = DATA / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload(owner, items), indent=2) + "\n", encoding="utf-8")


def re_sub_block(text: str, block: str) -> str:
    updated, count = re.subn(r'<section class="go-deeper"[\s\S]*?</section>', block, text, count=1)
    return updated if count else text


def mirror_sources(dest: Path, owner: str, items: list[dict]) -> None:
    """Write the per-page sources.json the visible block links to."""
    dest.parent.mkdir(parents=True, exist_ok=True)
    dest.write_text(json.dumps(payload(owner, items), indent=2) + "\n", encoding="utf-8")


def write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def add_docx_sources(document, items: list[dict]) -> None:
    document.add_heading("Sources and go deeper", level=1)
    document.add_paragraph(_intro())
    for item in items:
        label = KIND_LABEL[item["kind"]]
        checked = f"last checked {item['last_checked']}"
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{label}. ")
        if item["url"]:
            _hyperlink(paragraph, item["title"], item["url"])
            paragraph.add_run(f". {checked}")
        else:
            note = f" {item['note']}" if item["note"] else ""
            paragraph.add_run(f"{item['title']}. confirm.{note} {checked}")


def _hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1E3A8A")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(color)
    props.append(underline)
    run.append(props)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_pptx_sources(pres, items: list[dict]) -> None:
    chunks = [items[index : index + 8] for index in range(0, len(items), 8)] or [[]]
    for offset, chunk in enumerate(chunks):
        slide = pres.slides.add_slide(pres.slide_layouts[1])
        title = "Sources and go deeper" if offset == 0 else "Sources and go deeper, continued"
        slide.shapes.title.text = title
        frame = slide.placeholders[1].text_frame
        frame.clear()
        frame.word_wrap = True
        first = True
        if offset == 0:
            frame.text = _intro()
            first = False
        for item in chunk:
            paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
            first = False
            label = KIND_LABEL[item["kind"]]
            checked = f"last checked {item['last_checked']}"
            if item["url"]:
                run = paragraph.add_run()
                run.text = f"{label}. {item['title']}"
                run.font.size = Pt(14)
                run.hyperlink.address = item["url"]
                tail = paragraph.add_run()
                tail.text = f". {checked}"
                tail.font.size = Pt(14)
            else:
                note = f" {item['note']}" if item["note"] else ""
                run = paragraph.add_run()
                run.text = f"{label}. {item['title']}. confirm.{note} {checked}"
                run.font.size = Pt(14)
        if offset == len(chunks) - 1:
            legal = frame.add_paragraph()
            legal.text = "Independent free learning-only student project. Not affiliated or endorsed. Nothing is sold."
            for run in legal.runs:
                run.font.size = Pt(12)


def room_markdown(room: dict, items: list[dict], json_href: str) -> str:
    bullets = "\n".join(f"- {line}" for line in room["five"])
    body = [
        f"# {room['title']}",
        "",
        bullets,
        "",
        room["core"],
        "",
        room["example"],
        "",
        markdown_block(items, json_href),
    ]
    return "\n".join(body)


def sheet_markdown(title: str, lines: list[str], items: list[dict], json_href: str) -> str:
    bullets = "\n".join(f"- {line}" for line in lines)
    return f"# {title}\n\n{bullets}\n\n{markdown_block(items, json_href)}"


def append_pdf_sources(path: Path, items: list[dict]) -> None:
    """Add a final sources page to an existing PDF. Skip if one is already there."""
    reader = PdfReader(str(path))
    existing = "\n".join((page.extract_text() or "") for page in reader.pages)
    if "Sources and go deeper" in existing:
        return
    buffer_path = path.with_suffix(".sources-tmp.pdf")
    drawing = canvas.Canvas(str(buffer_path), pagesize=letter)
    width, height = letter
    y = height - 54
    drawing.setFont("Helvetica-Bold", 14)
    drawing.drawString(48, y, "Sources and go deeper")
    y -= 22
    drawing.setFont("Helvetica", 9)
    for line in plain_lines(items):
        text = line
        while text:
            chunk = text[:100]
            if len(text) > 100:
                cut = chunk.rfind(" ")
                if cut > 40:
                    chunk = chunk[:cut]
            if y < 54:
                drawing.showPage()
                drawing.setFont("Helvetica", 9)
                y = height - 54
            drawing.drawString(48, y, chunk.strip())
            y -= 12
            text = text[len(chunk) :].strip()
    drawing.save()
    extra = PdfReader(str(buffer_path))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    for page in extra.pages:
        writer.add_page(page)
    with path.open("wb") as handle:
        writer.write(handle)
    buffer_path.unlink(missing_ok=True)


def patch_flagship_cheats() -> int:
    changed = 0
    folder = ROOT / "flagship" / "cheats"
    for slug, keys in CHEAT_KEYS.items():
        path = folder / f"{slug}.html"
        if not path.is_file():
            continue
        items = entries(keys)
        rel = f"cheats/{slug}.json"
        write_json(rel, slug, items)
        mirror_sources(folder / "sources" / slug / "sources.json", slug, items)
        block = html_block(items, f"sources/{slug}/sources.json", dark=True)
        text = path.read_text(encoding="utf-8")
        needle = '<footer class="cheat-print">'
        if needle not in text:
            continue
        if "Sources and go deeper" in text:
            text = re_sub_block(text, block)
        else:
            text = text.replace(needle, block + needle, 1)
        path.write_text(text, encoding="utf-8")
        changed += 1
    return changed


def write_campus_bundle(items: list[dict]) -> None:
    block = html_block(items, "assets/sources.json")
    md = markdown_block(items, "assets/sources.json")
    payload_js = (
        "window.UC_GO_DEEPER_HTML = "
        + json.dumps(block)
        + ";\nwindow.UC_GO_DEEPER_MD = "
        + json.dumps(md)
        + ";\n"
    )
    write_text(ROOT / "campus" / "assets" / "go-deeper.js", payload_js)
    write_text(ROOT / "campus" / "assets" / "sources.json", json.dumps(payload("campus", items), indent=2) + "\n")
    write_json("campus.json", "campus", items)


def cheat_html(slug: str) -> str:
    keys = CHEAT_KEYS.get(slug, CORE)
    items = entries(keys)
    rel = f"cheats/{slug}.json"
    write_json(rel, slug, items)
    mirror_sources(ROOT / "flagship" / "cheats" / "sources" / slug / "sources.json", slug, items)
    return html_block(items, f"sources/{slug}/sources.json", dark=True)
